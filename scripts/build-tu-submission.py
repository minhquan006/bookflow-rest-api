#!/usr/bin/env python3
"""Generate Tu's BookFlow contribution report and editable presentation from reviewed content."""
from pathlib import Path
import json
import subprocess
import textwrap
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/submission/tu'
OUT.mkdir(parents=True, exist_ok=True)
GREEN = '#225e46'
SECTIONS = [
{'title':'BookFlow: Catalog and REST API Contract','subtitle':'Contribution chapter — Phạm Văn Tú, 52400252', 'paragraphs':[
'Topic: RESTful Web Services. Course: Web Programming & Applications (503073), Semester 1, Academic Year 2026–2027. Prepared on 22 October 2026.',
'This document covers Tú’s assigned contribution: a comparison of REST, GraphQL and SOAP; the Categories, Books and Copies module; API contracts; and the responsive Catalog interface. It is intended to be integrated into the group report, rather than submitted as the complete group report.',
'Nguyễn Minh Quân (52400227) owns Auth and Users. Phạm Văn Tú (52400252) owns Catalog and shared UI. Lê Nguyễn Vạn An (52400258) owns Loans, Dashboard/Audit and deployment/testing coordination. The existing Auth implementation is an integration prerequisite and is not reattributed to the Catalog contribution.',
'The implemented system is a React/TypeScript client, an Express/TypeScript JSON API and a PostgreSQL 17 database. Catalog functionality is connected to real database queries. Planned loan and account-management features are distinguished from implemented behavior throughout this chapter.',
'The Faculty template archive and instructor details have not been supplied. This contribution uses a consistent technical-report layout without inventing a university logo, faculty name or instructor. The final group report still needs the required cover pages, institutional formatting, 20–60 content pages and a recorded English presentation involving all members.'],
 'table':[['Deliverable','Status'],['Catalog implementation','Implemented and validated locally'],['Report, slides, narration script','Prepared for group integration'],['Formal Faculty layout / instructor','Requires group confirmation'],['Final three-member video with sound','Members must record their own presentation']]},
{'title':'1. Comparing three approaches on the same problem','paragraphs':[
'REST is an architectural style for distributed hypermedia systems, rather than a synonym for JSON endpoints [1]. A resource-oriented HTTP API can apply methods, representations and conditional requests to a library catalog. GraphQL specifies a typed query language and execution model: a client selects fields and traverses relationships through a schema [3]. SOAP defines an XML-based messaging framework with an envelope and processing rules; its binding framework is not limited to HTTP [4]. These are different kinds of specification, so the comparison should focus on a concrete application need.',
'For BookFlow, consider listing book metadata, opening one book and updating its title. REST expresses these as requests to book resources. GraphQL could expose a books query and an updateBook mutation. SOAP could describe application-specific operations such as GetBooks and UpdateBook using its message framework and a service contract. The project implements only the REST variant. The other two alternatives below are analytical designs, not deployed endpoints or measured implementations.',
'The principal selection criteria are contract clarity for a small team, integration with a browser client, authorization, caching semantics and tooling. Performance cannot be ranked without an equivalent implementation, workload and measurement method. No latency or throughput advantage is claimed in this report.'],
 'table':[['Criterion','REST over HTTP in BookFlow','GraphQL alternative','SOAP alternative'],['Abstraction','Resources and HTTP methods','Typed fields, queries, mutations','XML messages and operations'],['Returned fields','Defined by endpoint DTO','Selected by client schema query','Defined by message contract'],['Validation','Zod plus OpenAPI','Schema and resolver validation','Message/schema validation'],['Authorization','Route, principal and service rules','Resolver and business-layer rules','Service and message security rules'],['Caching','HTTP validators where appropriate','Operation/transport-aware design','Transport/application dependent'],['Primary cost here','More resource endpoints','Resolver and query-complexity work','XML/tooling and contract overhead']]},
{'title':'2. Equivalent use case and selection rationale','paragraphs':[
'The REST request below reads a bounded collection with a category filter and an allowlisted sort. Its response contains safe Book representations and pagination metadata. The client receives the endpoint-defined fields, including description, even if a particular card view only needs a title and author. This is a practical representation trade-off, not proof that REST always over-fetches.',
'A GraphQL alternative could let the card request select only title, author and category name in one typed operation. That flexibility would require a schema, resolver authorization and bounded query execution. It does not remove the need for database indexes, transactions or consistent business rules. A careless resolver implementation can still perform excessive database queries.',
'A SOAP alternative could define a GetBooks request and a corresponding XML response. SOAP can be appropriate where existing enterprise contracts, message processing and related standards are required. Those requirements are absent from this small library demonstration. BookFlow therefore selects resource-oriented HTTP with familiar status codes, OpenAPI and a shared browser HTTP client.',
'The HTTP representation is JSON for this project; REST itself does not require JSON. GraphQL is not inherently faster or more secure. SOAP is not restricted to HTTP. These qualifications avoid treating implementation conventions as universal properties of each approach.'],
 'code':'Implemented REST example:\nGET /api/v1/books?categoryId=<uuid>&page=1&limit=12&sort=title\n\nIllustrative GraphQL alternative (not implemented):\nquery { books(categoryId: "<uuid>", page: 1, limit: 12) {\n  items { id title author category { name } } total\n} }\n\nIllustrative SOAP operation (not implemented):\n<soap:Envelope xmlns:soap="http://www.w3.org/2003/05/soap-envelope">\n  <soap:Body><GetBooks><CategoryId>...</CategoryId></GetBooks></soap:Body>\n</soap:Envelope>'},
{'title':'3. Architecture and request flow','paragraphs':[
'The application follows controller → service → repository boundaries. The controller translates HTTP, validates input with strict Zod schemas and applies authentication/RBAC middleware. The service owns business rules and transaction boundaries. Repositories execute parameterized SQL on the client passed by the service. React receives camelCase DTOs through the existing shared HTTP client.',
'For example, a staff member submits a new book. The client sends JSON containing categoryId, isbn, title, author and description. The controller rejects unknown fields, verifies the access token and obtains the current role/status from PostgreSQL. The Catalog service verifies staff access and opens a transaction. The repository inserts metadata, and an audit entry is written on the same client. Only after commit does the controller return 201, Location and an ETag.',
'An invalid input returns a 400 Problem Details response. A member attempting a staff mutation receives 403. Unique-key conflicts are mapped to specific 409 codes rather than exposing PostgreSQL messages. If audit insertion fails, the transaction rolls back the mutation. The HTTP client does not automatically retry a mutation after a timeout, because the original operation may already have committed.'], 'image':'figures/architecture.png', 'caption':'Figure 1. Catalog request boundaries. Interactive transactions use the same PostgreSQL client.'},
{'title':'4. Resource model and validation','paragraphs':[
'A book is bibliographic metadata, while a copy is an individual physical item. This distinction allows one title to have several barcodes and conditions without duplicating title/author metadata. Categories classify books. Loans reference copies and must retain historical relationships. Consequently, deleting a book in the Catalog API means archiving metadata; changing a copy to RETIRED preserves its record.',
'Validation is applied before SQL execution. Names and titles are trimmed and length-bounded. Category names are unique case-insensitively; slugs follow lowercase kebab-case. A supplied ISBN must follow the agreed 10/13-character format, with X permitted at the end of ISBN-10. The current contract checks format, not the mathematical ISBN checksum. A null ISBN is permitted. Barcodes are trimmed, bounded and unique.',
'Search is a case-insensitive substring on title, author or ISBN. User values are SQL parameters; percent and underscore are escaped as literal search characters. Sort keys map to a fixed server allowlist and always add id as a deterministic tie-breaker. Collection count and rows use a repeatable-read snapshot, so pagination metadata agrees with the returned page within a request. Offset pagination may still shift between independent requests when other users change data.'],
 'table':[['Resource','Key fields','Consistency rule'],['Category','id, name, slug','Cannot delete while any book references it'],['Book','categoryId, isbn, title, author, description','version and archivedAt controlled by server'],['BookCopy','bookId, barcode, conditionStatus','New copies AVAILABLE; no condition change with an open loan'],['Audit','actor, action, entity, requestId','Commits or rolls back with the mutation']]},
{'title':'5. HTTP contract, permissions and errors','paragraphs':[
'The operational base path is /api/v1. GET reads representations; POST creates a category, book or copy; PUT replaces all writable book/category fields; PATCH updates an allowed subset; DELETE archives a book or removes an unreferenced category. A new resource returns 201 and Location. Successful deletion/archive returns 204 without a body. These behaviors illustrate the HTTP semantics discussed in RFC 9110 [2].',
'Public users can list categories and active books, read active book metadata and request availability. LIBRARIAN and ADMIN can mutate the catalog and list physical copies. MEMBER does not inherit staff permissions. Archived book metadata is hidden with 404 for public users; explicit includeArchived=true requires staff. Authentication and visibility are checked before a conditional GET can return 304.',
'Errors use application/problem+json and a requestId that can be correlated with logs [5]. The client displays a useful message and field errors where available. SQL details, password hashes and access tokens are not included in these error representations.'],
 'table':[['Route group','Methods','Result / significant failures'],['/categories','GET, POST','200 list; 201 create; 409 duplicate'],['/categories/{id}','PUT, DELETE','200 / 204; 409 CATEGORY_IN_USE'],['/books','GET, POST','200 list; 201 with Location/ETag'],['/books/{id}','GET','200 or 304; hidden/missing 404'],['/books/{id}','PUT, PATCH, DELETE','200 / 204; 428 missing validator; 412 stale'],['/books/{id}/availability','GET','200 fresh count; no-store'],['/books/{id}/copies','GET, POST','200 / 201; staff only'],['/copies/{id}','PATCH','200; 409 COPY_HAS_ACTIVE_LOAN']]},
{'title':'6. Concurrent edits, archiving and physical copies','paragraphs':[
'A book ETag is a strong metadata version such as "book-<uuid>-v2". Editing requires an exact If-Match value. The service locks the book row with FOR UPDATE, compares the version inside the transaction and then applies the update. Two writes using the same validator cannot silently overwrite one another: one succeeds and increments the version, while the other receives 412. A missing validator receives 428. This is optimistic concurrency at the HTTP boundary combined with database serialization for the critical section.',
'Archiving takes the book lock before checking whether any physical copy has an open loan. Changing a copy condition locks the copy before checking its open loan. These rules depend on the future Loans module following the shared book → copy locking convention when borrowing. The Catalog tests simulate a borrower with an independent database connection, observe the Catalog request waiting for the lock, commit a loan and verify a 409 response. They do not imply that the loan HTTP API has been implemented.',
'Availability is calculated from AVAILABLE copies without open loans. It is not part of the metadata version and always uses a fresh no-store operational response. A copy condition change can therefore alter availability without pretending the bibliographic metadata has changed.'],
 'table':[['Sequence','Session A','Session B'],['Read','GET returns ETag v1','GET returns ETag v1'],['First write','If-Match v1 succeeds → v2','Still editing an old representation'],['Conflicting write','Committed metadata remains','If-Match v1 → 412; reload required']]},
{'title':'7. Public Catalog interface','paragraphs':[
'The Catalog reuses shared Button, FormField, Feedback, Pagination and HTTP utilities instead of introducing separate request/error formats. A visitor can search, filter by category, choose an allowed sort, change pages and inspect metadata with the current availability. The UI shows actual API responses rather than placeholder book data.',
'Loading, empty and error states are explicit. AbortController cancels requests when a screen or filter changes. A loader identity prevents late responses from replacing newer results. Search/filter changes reset the page where appropriate. Server-side permissions remain authoritative even when management controls are not rendered for a visitor.'],
 'image':'figures/public-catalog.png', 'caption':'Figure 2. Real Catalog capture from an isolated test schema. Synthetic fixture data; no mock API.'},
{'title':'8. Staff workflow and conflict recovery','paragraphs':[
'Staff sign-in uses the existing Auth API. The access token is kept in memory, not localStorage. A refresh of the page requires another sign-in because the refresh/logout endpoints are still planned. The Sign out action currently removes the local access token; it does not revoke a server-side refresh family.',
'The staff interface can create/edit categories, add/edit/archive books, create copies and change physical conditions. Deletion/archive require confirmation. Buttons block repeated submissions while work is pending, and database uniqueness/transaction rules enforce correctness independently of the UI. A failed mutation preserves the form so the user can correct it.',
'The conflict example is produced by a real second API request after the edit form has loaded its validator. Submitting that stale form receives 412. The UI asks the user to reload the current book before trying again; it does not silently resubmit against a new version.'],
 'image':'figures/stale-version.png', 'caption':'Figure 3. Actual stale-version recovery in the Catalog form. The screenshot is a real API response, not a visual mockup.'},
{'title':'9. Validation and evidence','paragraphs':[
'The development environment uses Node 24.19.0, npm 11 and PostgreSQL 17. Tests are run against the repository commands, and only executed results are reported. Integration tests require a database name ending _test and create independent random schemas that are cleaned up afterwards. Browser Catalog tests run a separate API scoped to a test schema, so staff mutations do not modify the development database.',
'The current validation contains 59 foundation unit/HTTP/component tests, 52 integration tests and 10 browser tests: 121 passed. Of the integration tests, 24 are Catalog-specific and 28 cover Auth/database foundations. Four browser cases cover Catalog in desktop/mobile projects; the other six are foundation regressions. These numbers are local results, not an assertion that GitHub CI or production deployment has passed.',
'Meaningful Catalog cases include guest/member rejection, invalid UUIDs/body/query data, duplicate category/ISBN/barcode conflicts, pagination/search/sort, ETag 304/412/428, concurrent same-version writes, archived visibility, loan-related archive/condition guards and rollback when audit insertion fails. Browser cases exercise creation/editing/archive, copy availability, conflict reload, category deletion errors, empty/error/retry states and responsive overflow checks.'],
 'table':[['Executed command','Outcome','Evidence'],['npm run check','Passed; includes 59 tests','logs/check.txt'],['npm run test:integration','52 passed, 0 failed','logs/integration.txt'],['npm run test:e2e','10 passed, desktop/mobile','logs/e2e.txt'],['Visual rehearsal recording','Real isolated browser/API workflow','video/ + logs/recording-http.json']]},
{'title':'10. Limitations, integration and conclusion','paragraphs':[
'The Catalog contribution demonstrates resource modeling, method/status semantics, strict contracts, authorization, conditional requests and transaction-protected rules through a connected UI. It does not establish full compliance with every REST constraint: the API has limited hypermedia, and most responses are deliberately no-store. The project should therefore explain its resource-oriented HTTP design accurately instead of claiming that JSON routes alone prove complete REST maturity.',
'Known limitations include offset pagination for deep pages, a single category and author string per book, format-only ISBN checks, and a small synthetic data set. No performance benchmark, high-availability deployment or second-machine restoration result is fabricated. Front-end accessibility work covers labels, keyboard-friendly native controls, visible feedback and responsive checks; it is not a formal accessibility certification.',
'The next team integration work is Auth refresh/logout and full account screens, plus Loans, Dashboard and Audit UI. Loan service locking must remain consistent with the Catalog safeguards. A draft PR should be reviewed against the existing feature branches before merge; one person should not overwrite another member’s history to make a contribution appear complete.',
'Tú’s presentation can now explain a full Catalog request flow, the REST/GraphQL/SOAP choice and a concurrent-edit failure with evidence. The final submission remains a group responsibility: apply the Faculty template, verify instructor/course metadata, combine chapters into 20–60 content pages and record all three members speaking English within the 20-minute limit.']},
{'title':'References and artifact guide','paragraphs':[
'[1] R. T. Fielding, Architectural Styles and the Design of Network-based Software Architectures, doctoral dissertation, 2000, Chapter 5. https://ics.uci.edu/~fielding/pubs/dissertation/rest_arch_style.htm',
'[2] IETF, RFC 9110: HTTP Semantics, 2022. In particular, method semantics, conditional requests and response status codes. https://www.rfc-editor.org/rfc/rfc9110.html',
'[3] GraphQL Foundation, GraphQL Specification, September 2025 edition. Type system, validation and execution. https://spec.graphql.org/September2025/',
'[4] W3C, SOAP Version 1.2 Part 1: Messaging Framework, Second Edition, 2007. https://www.w3.org/TR/soap12-part1/',
'[5] IETF, RFC 9457: Problem Details for HTTP APIs, 2023. https://www.rfc-editor.org/rfc/rfc9457.html',
'[6] node-postgres documentation, Transactions. https://node-postgres.com/features/transactions',
'[7] Course assignment: 503073-Midterm-V1.docx, sections III–V, stored in docs/. The separate Faculty template archive is not present in the repository.',
'Companion artifacts: BookFlow_Tu_Slides.pptx is editable and contains speaker notes; BookFlow_Tu_Slides.pdf is a reading copy. SPEAKER_SCRIPT.md provides an approximately six-minute English segment. The silent practice MP4 and subtitles are rehearsal aids and must not be submitted as the required group video with sound.',
'Full implementation/route details: docs/CATALOG_HANDOFF.md and docs/openapi.yaml. Existing real seed-data screenshots are in docs/evidence/. This chapter’s figures come from a fresh isolated test fixture. The references above identify authoritative standards; direct GraphQL/W3C retrieval was blocked by this environment’s network policy, so no claim of a live standards-site verification is made.']}
]
SLIDES = [
('BookFlow: Catalog and API Contract',['Phạm Văn Tú · 52400252','REST / GraphQL / SOAP comparison','A connected Catalog with safe concurrent edits'],None,25,
"My part of BookFlow is the Catalog and its API contract. I will explain why we chose a resource-oriented HTTP API, how Categories, Books and Copies are modeled, and how the UI handles real failures. The implementation uses React, Express and PostgreSQL. The evidence shown here is from executed local checks, rather than claims about features that are still planned."),
('Model the physical library, not just a list',['Category classifies a title','Book stores metadata; Copy stores a unique barcode','Archiving preserves history'], 'figures/architecture.png',30,
"A title is different from an individual physical copy. Several copies can share book metadata, while each copy has its own barcode and condition. This separation lets us calculate availability without duplicating title information. Requests pass through a controller, a service and a repository. The service owns the transaction and writes the audit entry using the same database client as the mutation."),
('REST, GraphQL and SOAP: compare the same task',['REST: resources, methods and HTTP validators','GraphQL: typed field selection; resolver and query limits','SOAP: XML messaging and operation contracts','No performance ranking without an equivalent benchmark'],None,55,
"These approaches are not identical kinds of standard. REST is an architectural style. GraphQL specifies a typed query language and execution model. SOAP specifies an XML messaging framework. For a book card, our REST endpoint returns a defined book representation. A GraphQL design could let the client request only title, author and category name, but would still need authorization and bounded query execution. A SOAP design could define a GetBooks operation and its XML contract. We chose HTTP resources because the team can demonstrate clear methods, statuses, validation and OpenAPI tooling. This is a qualitative comparison: we did not build equivalent GraphQL and SOAP servers or measure their speed. REST does not require JSON, and SOAP is not limited to HTTP."),
('A contract the browser can use',['GET reads; POST creates; PUT replaces; PATCH changes a subset','201 + Location for creation; 204 without a response body','400 invalid · 403 forbidden · 409 conflict','OpenAPI + Postman describe implemented operations'],None,40,
"The contract is more than a list of URLs. Creating a book requires all writable metadata, rejects unknown fields and returns status 201 with a Location and an ETag. PUT requires the complete editable representation, while PATCH accepts an allowed nonempty subset. Deleting a book archives it and returns 204. Errors use Problem Details with a request ID, so the client and logs can refer to the same failure. OpenAPI distinguishes implemented endpoints from planned ones, and the Postman collection provides concrete request examples."),
('Public UI and staff access',['Search title / author / ISBN','Filter, allowlisted sort and bounded pagination','Public reads; LIBRARIAN/ADMIN mutations','Server permissions are authoritative'], 'figures/public-catalog.png',35,
"A visitor can search by title, author or ISBN, filter by category and choose a sort order. The server maps sort names to fixed SQL columns, and all user values are parameters. The page shows loading, empty and retry states. Staff sign-in exposes management forms, but hiding buttons is only a user-interface choice: the API independently verifies the token and reads the current role and status from the database. A member calling a staff mutation receives 403."),
('Demonstrate a real management flow',['Create a category and book','Add a physical copy; change condition','Availability changes independently of metadata','Duplicate barcode returns 409'], 'figures/physical-copies.png',50,
"In the live Catalog, I first create a category and then a book that references it. Next I add a copy with a unique barcode. The initial condition is AVAILABLE. Changing it to MAINTENANCE reduces the available count; restoring it increases that count. Availability also excludes copies with an open loan. If I try to reuse the barcode, the server returns 409 and the UI keeps the error visible. These are real HTTP and database operations. The rehearsal uses an isolated test schema, so it does not change our development data."),
('Two editors: detect a stale version',['GET returns a strong metadata ETag','Mutation sends If-Match','One writer advances v1 → v2','A stale writer receives 412 and reloads'], 'figures/stale-version.png',45,
"A book’s metadata version becomes a strong ETag. When a form saves, it sends that value in If-Match. Inside a transaction, the service locks the book row, checks the version and updates the metadata. If another request has already saved, the old form receives 412 instead of overwriting the newer data. Missing If-Match receives 428. I demonstrate this by changing the book through a second real API request while the first form is open. The UI then asks me to reload before saving again."),
('Protect history and open loans',['Archive a title; do not erase its history','Referenced categories cannot be deleted','Open loans block archive and copy-condition changes','Audit commits with the mutation'],None,35,
"Archiving hides a book from public access while retaining its copies and historical relationships. A category referenced by an active or archived book cannot be deleted. An open loan blocks archiving and changing the condition of its copy. The Catalog locks the relevant book or copy before checking loans. The future Loans service must follow the same lock order. Our database tests simulate a competing borrower and verify that the Catalog waits, then rejects an unsafe change. Audit failures also roll back the mutation."),
('Evidence: separate local checks from CI',['59 foundation tests + 52 integration + 10 browser = 121 passed','24 Catalog integration cases','Desktop/mobile flows, errors and conflict recovery','No invented performance or production claims'],None,30,
"The recorded validation has 121 passing tests: 59 foundation tests, 52 integration tests and 10 browser tests. Twenty-four integration cases specifically cover the Catalog. The browser cases run on desktop and mobile and exercise real creation, editing, copy conditions, archive and error recovery. Integration tests use independent schemas in a test database. These are local results. I do not describe them as GitHub CI, a production deployment, a performance benchmark or proof that the Loans module has been completed."),
('Integration and remaining group work',['Catalog is implemented; loan flows remain separate','Refresh/logout and full account UI remain planned','Combine chapters and apply the Faculty template','Record all three members in English; total ≤20 minutes'],None,15,
"This completes the Catalog contribution, not the entire group submission. We still need to integrate the other modules, apply the official Faculty template and record all three members speaking English within twenty minutes. The editable slides, report chapter, source and test evidence are ready for that group review.")
]

# Use fonts embedded in PDFs; the Word file remains editable.
FONT = Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('Body','DejaVuSerif.ttf'),('BodyBold','DejaVuSerif-Bold.ttf'),('Sans','DejaVuSans.ttf'),('SansBold','DejaVuSans-Bold.ttf'),('Mono','DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(FONT/file)))
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='ReportBody',fontName='Body',fontSize=10.5,leading=15,spaceAfter=10))
styles.add(ParagraphStyle(name='ReportTitle',fontName='SansBold',fontSize=22,leading=28,textColor=colors.HexColor(GREEN),spaceAfter=18))
styles.add(ParagraphStyle(name='CaptionSmall',fontName='Sans',fontSize=8.5,leading=11,textColor=colors.HexColor('#566c62'),spaceAfter=10))
styles.add(ParagraphStyle(name='CellSmall',fontName='Sans',fontSize=8,leading=10))
styles.add(ParagraphStyle(name='CellHead',fontName='SansBold',fontSize=8,leading=10,textColor=colors.white))
doc=Document(); sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('BookFlow · Tú contribution · ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
story=[];md=[]
for i,s in enumerate(SECTIONS):
    if i: doc.add_page_break();story.append(PageBreak())
    doc.add_heading(s['title'],0 if i==0 else 1);story.append(Paragraph(escape(s['title']),styles['ReportTitle']));md.append('# '+s['title']+'\n')
    if s.get('subtitle'):
        doc.add_paragraph(s['subtitle'],'Subtitle');story.append(Paragraph(escape(s['subtitle']),styles['ReportBody']));md.append(s['subtitle']+'\n')
    for p in s['paragraphs']:
        doc.add_paragraph(p);story.append(Paragraph(escape(p),styles['ReportBody']));md.append(p+'\n')
    if s.get('table'):
        rows=s['table']; table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for c,t in zip(table.rows[0].cells,rows[0]):c.text=t
        for row in rows[1:]:
            for c,t in zip(table.add_row().cells,row):c.text=t
        data=[[Paragraph(escape(t),styles['CellHead'] if r==0 else styles['CellSmall']) for t in row] for r,row in enumerate(rows)]
        t=Table(data,colWidths=[(A4[0]-88)/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(GREEN)),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#dbe4db')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f6f3')])]))
        story.extend([Spacer(1,7),t]);md.extend(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(row)+' |' for row in rows[1:]]+[''])
    if s.get('code'):
        p=doc.add_paragraph();run=p.add_run(s['code']);run.font.name='Courier New';run.font.size=Pt(9)
        story.append(Preformatted(s['code'],ParagraphStyle('code',fontName='Mono',fontSize=7.4,leading=11),maxLineLength=91));md.append('```text\n'+s['code']+'\n```\n')
    if s.get('image'):
        path=OUT/s['image'];w,h=PILImage.open(path).size; maxw=A4[0]-88;maxh=290 if i==3 else 300;scale=min(maxw/w,maxh/h)
        story.extend([Spacer(1,8),Image(str(path),width=w*scale,height=h*scale),Paragraph(escape(s['caption']),styles['CaptionSmall'])])
        doc.add_picture(str(path),width=Inches(min(6.3,w*scale/72)));doc.add_paragraph(s['caption'],'Caption');md.append(f"![{s['caption']}]({s['image']})\n")
    story.append(Spacer(1,6))
doc.save(OUT/'BookFlow_Tu_Report.docx')
def footer_pdf(c,d):
    c.setFont('Sans',8);c.setFillColor(colors.HexColor('#566c62'));c.drawString(44,25,'BookFlow · Phạm Văn Tú · Contribution chapter');c.drawRightString(A4[0]-44,25,str(d.page))
SimpleDocTemplate(str(OUT/'BookFlow_Tu_Report.pdf'),pagesize=A4,leftMargin=44,rightMargin=44,topMargin=42,bottomMargin=45,title='BookFlow Catalog — Tu contribution',author='Phạm Văn Tú').build(story,onFirstPage=footer_pdf,onLaterPages=footer_pdf)
(OUT/'BookFlow_Tu_Report.md').write_text('\n'.join(md),encoding='utf8')

prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
script=['# Tú — English speaking script (~6 minutes)\n','Target segment: six minutes. Rehearse at approximately 125–145 words per minute; timings are targets, not a claim about a recorded voice.\n']
slidepdf=canvas.Canvas(str(OUT/'BookFlow_Tu_Slides.pdf'),pagesize=(960,540));slidepdf.setTitle('BookFlow — Tu Catalog contribution')
for idx,(title,bullets,image,seconds,notes) in enumerate(SLIDES,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F4F6F3')
    shape=slide.shapes.add_shape(1,0,0,prs.slide_width,PI(.13));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('225E46');shape.line.fill.background()
    box=slide.shapes.add_textbox(PI(.6),PI(.4),PI(12),PI(.9));tf=box.text_frame;tf.word_wrap=True;p=tf.paragraphs[0];p.text=title;p.font.size=PP(30);p.font.bold=True;p.font.color.rgb=RGBColor.from_string('19352E')
    text_width=5.8 if image else 11.8
    box=slide.shapes.add_textbox(PI(.7),PI(1.65),PI(text_width),PI(4.7));tf=box.text_frame;tf.word_wrap=True
    for j,bullet in enumerate(bullets):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text='• '+bullet;p.font.size=PP(22 if image else 25);p.font.color.rgb=RGBColor.from_string('19352E');p.space_after=PP(22)
    if image:
        path=OUT/image;w,h=PILImage.open(path).size;scale=min(5.6/w,4.7/h);slide.shapes.add_picture(str(path),PI(7),PI(1.65),width=PI(w*scale),height=PI(h*scale))
    box=slide.shapes.add_textbox(PI(.7),PI(6.95),PI(12),PI(.3));p=box.text_frame.paragraphs[0];p.text=f'BookFlow · Phạm Văn Tú · 52400252                              {idx}/10';p.font.size=PP(11)
    slide.notes_slide.notes_text_frame.text=f'Target: {seconds} seconds\n\n{notes}'
    script.extend([f'## Slide {idx}: {title} — target {seconds}s\n',notes+'\n'])
    slidepdf.setFillColor(colors.HexColor('#f4f6f3'));slidepdf.rect(0,0,960,540,fill=1,stroke=0);slidepdf.setFillColor(colors.HexColor(GREEN));slidepdf.rect(0,531,960,9,fill=1,stroke=0)
    y=494;slidepdf.setFont('SansBold',24);slidepdf.setFillColor(colors.HexColor('#19352e'))
    for line in textwrap.wrap(title,width=57):slidepdf.drawString(42,y,line);y-=29
    y=405;slidepdf.setFont('Sans',17 if image else 20)
    for bullet in bullets:
        for line in textwrap.wrap('• '+bullet,width=43 if image else 79):slidepdf.drawString(48,y,line);y-=25 if image else 29
        y-=15
    if image:
        path=OUT/image;w,h=PILImage.open(path).size;scale=min(405/w,335/h);slidepdf.drawImage(str(path),515,100+(335-h*scale)/2,width=w*scale,height=h*scale,mask='auto')
    slidepdf.setFont('Sans',9);slidepdf.drawString(44,22,f'BookFlow · Phạm Văn Tú · 52400252');slidepdf.drawRightString(917,22,f'{idx}/10');slidepdf.showPage()
prs.save(OUT/'BookFlow_Tu_Slides.pptx');slidepdf.save();(OUT/'SPEAKER_SCRIPT.md').write_text('\n'.join(script),encoding='utf8')
(OUT/'requirements.txt').write_text('python-docx==1.2.0\npython-pptx==1.0.2\nreportlab==4.4.9\nPillow==12.3.0\n')
(OUT/'content.json').write_text(json.dumps({'report':SECTIONS,'slides':[{'title':s[0],'bullets':s[1],'image':s[2],'seconds':s[3],'notes':s[4]} for s in SLIDES]},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print('Created Word/PDF report, editable slides + PDF, source content and English speaker script.')
print('Target speaking seconds:',sum(s[3] for s in SLIDES),'words:',sum(len(s[4].split()) for s in SLIDES))

subprocess.run([str(ROOT / "node_modules/.bin/prettier"), "--write", str(OUT / "BookFlow_Tu_Report.md"), str(OUT / "SPEAKER_SCRIPT.md"), str(OUT / "content.json")], cwd=ROOT, check=True)
